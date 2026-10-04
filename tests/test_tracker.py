from pathlib import Path

from apply_bot.filter.legitimacy import FilterResult
from apply_bot.models import JobPost
from apply_bot.pipeline import job_from_row
from apply_bot.store import db


def test_careers_click_without_thanks_stays_opened():
    assert db.apply_status_for("careers_opened", reason="extra_step") == (
        "Careers opened (login or extra step)"
    )
    assert db.apply_status_for("form_submitted") == "Form submitted"


def test_record_keeps_stipend_email_and_url(tmp_path: Path):
    conn = db.connect(tmp_path / "apps.db")
    post = JobPost(
        company="RoadVision AI",
        role="SDE Intern",
        stipend="₹25,000/month",
        location="New Delhi",
        batch_raw="2025/2026/2027",
        apply_email="hr@roadvision.ai",
        apply_url="https://forms.gle/example",
        apply_method="email_and_form",
        source_fingerprint="rv1",
    )
    result = FilterResult(verdict="SEND", reason="company_domain")
    db.record(conn, post, result, "draft_created")
    csv_path = tmp_path / "job_tracker.csv"
    xlsx_path = db.export_tracker(conn, csv_path)
    text = csv_path.read_text(encoding="utf-8")
    assert "RoadVision AI" in text
    assert "SDE Intern" in text
    assert "25,000" in text
    assert "hr@roadvision.ai" in text
    assert "forms.gle/example" in text
    assert "draft_created" in text
    assert xlsx_path.suffix == ".xlsx"
    assert xlsx_path.exists()
    from openpyxl import load_workbook

    book = load_workbook(xlsx_path)
    assert "All jobs" in book.sheetnames
    assert "Applied" in book.sheetnames
    assert book["Applied"]["B2"].value == "RoadVision AI"
    assert book["Applied"]["M2"].value == "Draft (not sent)"
    assert "Draft (not sent)" in text
    row = conn.execute("SELECT * FROM applications WHERE fingerprint = 'rv1'").fetchone()
    assert row["stipend"] == "₹25,000/month"
    assert row["apply_email"] == "hr@roadvision.ai"
    assert row["apply_url"] == "https://forms.gle/example"
    assert row["apply_status"] == "Draft (not sent)"
    assert db.fingerprint_seen(conn, "rv1") is False
    assert db.already_applied(conn, "RoadVision AI", "SDE Intern") is False
    db.record(conn, post, result, "form_submitted")
    assert db.fingerprint_seen(conn, "rv1") is True
    assert db.already_applied(conn, "RoadVision AI", "SDE Intern") is True
    conn.close()


def test_record_round_trips_raw_text(tmp_path: Path):
    conn = db.connect(tmp_path / "apps.db")
    post = JobPost(
        company="Pazy",
        role="Fullstack Intern",
        batch_raw="2027",
        apply_url="https://example.com/form",
        apply_method="google_form",
        raw_text="Company - Pazy\nRole - Fullstack Intern\nReact Node batch 2027",
        source_fingerprint="pazy1",
    )
    result = FilterResult(verdict="REVIEW", reason="form_needs_click")
    db.record(conn, post, result, "form_prepared")
    row = conn.execute("SELECT * FROM applications WHERE fingerprint = 'pazy1'").fetchone()
    assert "React Node" in (row["raw_text"] or "")
    csv_path = tmp_path / "job_tracker.csv"
    db.export_tracker(conn, csv_path)
    assert "React Node" not in csv_path.read_text(encoding="utf-8")
    restored = job_from_row(row)
    assert "React Node" in restored.raw_text
    conn.close()
