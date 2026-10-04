from __future__ import annotations

import csv
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from apply_bot.models import FilterResult, JobPost


SCHEMA = """
CREATE TABLE IF NOT EXISTS applications (
    id INTEGER PRIMARY KEY,
    created_at TEXT NOT NULL,
    fingerprint TEXT NOT NULL UNIQUE,
    company TEXT,
    role TEXT,
    verdict TEXT,
    reason TEXT,
    apply_contact TEXT,
    status TEXT NOT NULL,
    screenshot_path TEXT,
    stipend TEXT,
    location TEXT,
    batch TEXT,
    apply_email TEXT,
    apply_url TEXT,
    apply_method TEXT,
    apply_status TEXT,
    raw_text TEXT
);
CREATE INDEX IF NOT EXISTS idx_company_role
    ON applications (company, role);
"""

EXTRA_COLUMNS = (
    ("stipend", "TEXT"),
    ("location", "TEXT"),
    ("batch", "TEXT"),
    ("apply_email", "TEXT"),
    ("apply_url", "TEXT"),
    ("apply_method", "TEXT"),
    ("apply_status", "TEXT"),
    ("raw_text", "TEXT"),
)

TRACKER_FIELDS = (
    "created_at",
    "company",
    "role",
    "stipend",
    "location",
    "batch",
    "apply_email",
    "apply_url",
    "apply_method",
    "verdict",
    "reason",
    "status",
    "apply_status",
)

BOT_TO_APPLY_STATUS = {
    "draft_created": "Draft (not sent)",
    "form_prepared": "Form filled (not submitted)",
    "careers_opened": "Careers opened (not submitted)",
    "needs_login": "Careers opened (login or extra step)",
    "email_sent": "Sent",
    "form_submitted": "Form submitted",
    "skipped": "Skipped",
    "seen": "Pending",
}

FINAL_APPLY_STATUSES = {"Sent", "Form submitted"}


def connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.executescript(SCHEMA)
    _migrate(conn)
    _backfill_contacts(conn)
    _backfill_apply_status(conn)
    return conn


def _migrate(conn: sqlite3.Connection) -> None:
    existing = {row[1] for row in conn.execute("PRAGMA table_info(applications)")}
    for name, typ in EXTRA_COLUMNS:
        if name not in existing:
            conn.execute(f"ALTER TABLE applications ADD COLUMN {name} {typ}")
    conn.commit()


def _backfill_contacts(conn: sqlite3.Connection) -> None:
    rows = conn.execute(
        "SELECT id, apply_contact, apply_email, apply_url FROM applications"
    ).fetchall()
    for row in rows:
        if row["apply_email"] or row["apply_url"] or not row["apply_contact"]:
            continue
        contact = row["apply_contact"].strip()
        if contact.startswith("http"):
            conn.execute(
                "UPDATE applications SET apply_url = ? WHERE id = ?",
                (contact, row["id"]),
            )
        elif "@" in contact:
            conn.execute(
                "UPDATE applications SET apply_email = ? WHERE id = ?",
                (contact, row["id"]),
            )
    conn.commit()


def apply_status_for(bot_status: str, existing: str | None = None, reason: str | None = None) -> str:
    if existing in FINAL_APPLY_STATUSES:
        return existing
    if bot_status == "skipped" and reason in {
        "stuck_ats",
        "skip_ragaai",
        "needs_login",
        "extra_step",
        "stuck_picker",
    }:
        return f"Skipped (stuck) — {reason}"
    if bot_status == "careers_opened" and reason in {"needs_login", "extra_step"}:
        return "Careers opened (login or extra step)"
    if bot_status == "form_prepared" and reason and (
        reason.startswith("missing:")
        or reason in {"submit_not_confirmed", "needs_login", "form_closed", "job_closed"}
    ):
        return f"Form filled (not submitted) — {reason}"
    return BOT_TO_APPLY_STATUS.get(bot_status, bot_status or "Pending")


def _backfill_apply_status(conn: sqlite3.Connection) -> None:
    rows = conn.execute("SELECT id, status, apply_status FROM applications").fetchall()
    for row in rows:
        if row["apply_status"] in FINAL_APPLY_STATUSES:
            continue
        mapped = apply_status_for(row["status"] or "", row["apply_status"])
        if mapped != (row["apply_status"] or ""):
            conn.execute(
                "UPDATE applications SET apply_status = ? WHERE id = ?",
                (mapped, row["id"]),
            )
    conn.commit()


def _norm(value: str | None) -> str:
    return (value or "").strip().lower()


DONE_STATUSES = frozenset({"email_sent", "form_submitted"})


def done_count(conn: sqlite3.Connection) -> int:
    row = conn.execute(
        "SELECT COUNT(*) AS n FROM applications WHERE status IN ('email_sent', 'form_submitted')"
    ).fetchone()
    return int(row["n"] if row else 0)
INCOMPLETE_STATUSES = frozenset(
    {"seen", "form_prepared", "careers_opened", "draft_created"}
)


def fingerprint_seen(conn: sqlite3.Connection, fingerprint: str) -> bool:
    row = conn.execute(
        "SELECT status, reason, company, role, apply_url FROM applications WHERE fingerprint = ?",
        (fingerprint,),
    ).fetchone()
    if not row:
        return False
    if row["status"] in DONE_STATUSES:
        return True
    from apply_bot.stuck import stuck_reason

    return bool(
        row["status"] == "skipped"
        and stuck_reason(row["company"], row["role"], row["apply_url"], row["reason"])
    )


def already_applied(conn: sqlite3.Connection, company: str | None, role: str | None) -> bool:
    if not company or not role:
        return False
    row = conn.execute(
        """
        SELECT 1 FROM applications
        WHERE lower(trim(company)) = ? AND lower(trim(role)) = ?
          AND status IN ('email_sent', 'form_submitted')
        """,
        (_norm(company), _norm(role)),
    ).fetchone()
    return row is not None


def incomplete_rows(conn: sqlite3.Connection) -> list[sqlite3.Row]:
    from apply_bot.stuck import completable_method, stuck_reason

    rows = conn.execute(
        """
        SELECT fingerprint, company, role, stipend, location, batch,
               apply_email, apply_url, apply_method, status, raw_text, reason
        FROM applications
        WHERE status IN ('seen', 'form_prepared', 'careers_opened', 'draft_created')
        ORDER BY company, role
        """
    ).fetchall()
    return [
        row
        for row in rows
        if completable_method(row["apply_method"])
        and not stuck_reason(row["company"], row["role"], row["apply_url"], row["reason"])
    ]


def mark_stuck_skipped(conn: sqlite3.Connection) -> list[tuple[str, str, str]]:
    from apply_bot.stuck import stuck_reason

    rows = conn.execute(
        """
        SELECT fingerprint, company, role, apply_url, reason, status
        FROM applications
        WHERE status IN ('seen', 'form_prepared', 'careers_opened', 'draft_created')
        """
    ).fetchall()
    marked: list[tuple[str, str, str]] = []
    for row in rows:
        why = stuck_reason(row["company"], row["role"], row["apply_url"], row["reason"])
        if not why:
            continue
        conn.execute(
            """
            UPDATE applications
            SET status = 'skipped', reason = ?, apply_status = ?
            WHERE fingerprint = ?
            """,
            (why, f"Skipped (stuck) — {why}", row["fingerprint"]),
        )
        marked.append((row["company"] or "?", row["role"] or "?", why))
    conn.commit()
    return marked


def record(
    conn: sqlite3.Connection,
    post: JobPost,
    result: FilterResult | None,
    status: str,
    screenshot_path: str | None = None,
) -> None:
    contact = post.apply_email or post.apply_url or ""
    existing = conn.execute(
        "SELECT apply_status FROM applications WHERE fingerprint = ?",
        (post.source_fingerprint,),
    ).fetchone()
    apply_status = apply_status_for(
        status,
        existing["apply_status"] if existing else None,
        result.reason if result else None,
    )
    conn.execute(
        """
        INSERT INTO applications
        (created_at, fingerprint, company, role, verdict, reason, apply_contact, status,
         screenshot_path, stipend, location, batch, apply_email, apply_url, apply_method,
         apply_status, raw_text)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(fingerprint) DO UPDATE SET
            company = excluded.company,
            role = excluded.role,
            verdict = excluded.verdict,
            reason = excluded.reason,
            apply_contact = excluded.apply_contact,
            status = excluded.status,
            screenshot_path = excluded.screenshot_path,
            stipend = excluded.stipend,
            location = excluded.location,
            batch = excluded.batch,
            apply_email = excluded.apply_email,
            apply_url = excluded.apply_url,
            apply_method = excluded.apply_method,
            apply_status = CASE
                WHEN applications.apply_status IN ('Sent', 'Form submitted')
                THEN applications.apply_status
                ELSE excluded.apply_status
            END,
            raw_text = COALESCE(excluded.raw_text, applications.raw_text)
        """,
        (
            datetime.now(timezone.utc).isoformat(),
            post.source_fingerprint,
            post.company,
            post.role,
            result.verdict if result else None,
            result.reason if result else None,
            contact,
            status,
            screenshot_path,
            post.stipend,
            post.location,
            post.batch_raw,
            post.apply_email,
            post.apply_url,
            post.apply_method,
            apply_status,
            post.raw_text or None,
        ),
    )
    conn.commit()


def _row_values(row: sqlite3.Row) -> list[str]:
    return [row[field] or "" for field in TRACKER_FIELDS]


def export_tracker(conn: sqlite3.Connection, csv_path: Path) -> Path:
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    rows = conn.execute(
        f"""
        SELECT {", ".join(TRACKER_FIELDS)}
        FROM applications
        ORDER BY created_at DESC, id DESC
        """
    ).fetchall()
    with csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(TRACKER_FIELDS))
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row[field] or "" for field in TRACKER_FIELDS})
    xlsx_path = csv_path.with_suffix(".xlsx")
    _write_excel(xlsx_path, rows)
    return xlsx_path


def _write_excel(xlsx_path: Path, rows: list[sqlite3.Row]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = [
        "Date",
        "Company",
        "Role / Profile",
        "Stipend",
        "Location",
        "Batch",
        "Apply email",
        "Apply link",
        "Apply method",
        "Verdict",
        "Reason",
        "Bot action",
        "Apply status",
    ]
    applied_status = {
        "draft_created",
        "form_prepared",
        "careers_opened",
        "email_sent",
        "form_submitted",
    }
    wb = Workbook()
    all_sheet = wb.active
    all_sheet.title = "All jobs"
    applied_sheet = wb.create_sheet("Applied")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def fill_sheet(sheet, selected) -> None:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in selected:
            sheet.append(_row_values(row))
        widths = [22, 32, 36, 22, 18, 22, 36, 50, 16, 12, 22, 18, 28]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"

    fill_sheet(all_sheet, rows)
    fill_sheet(applied_sheet, [row for row in rows if (row["status"] or "") in applied_status])
    wb.save(xlsx_path)


def tracker_path(project_root: Path) -> Path:
    return project_root / "data" / "job_tracker.csv"


def tracker_xlsx_path(project_root: Path) -> Path:
    return project_root / "data" / "job_tracker.xlsx"


def mark_apply_status(conn: sqlite3.Connection, company: str, apply_status: str, role: str | None = None) -> int:
    if role:
        cursor = conn.execute(
            """
            UPDATE applications
            SET apply_status = ?
            WHERE lower(trim(company)) = ? AND lower(trim(role)) = ?
            """,
            (apply_status, _norm(company), _norm(role)),
        )
    else:
        cursor = conn.execute(
            """
            UPDATE applications
            SET apply_status = ?
            WHERE lower(trim(company)) = ?
            """,
            (apply_status, _norm(company)),
        )
    conn.commit()
    return cursor.rowcount
